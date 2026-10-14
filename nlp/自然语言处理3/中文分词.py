# 导入核心库：docx文档读取、中文分词、词云生成、文本处理、文件操作
import jieba
import wordcloud
import re
from collections import Counter
from docx import Document
import os

# -------------------------- 1. 读取源文档：提取《最优化算法》教学大纲文本 --------------------------
def read_syllabus_docx(doc_path):
    """
    读取.docx格式大纲，提取段落和表格文本（去重处理重复内容）
    :param doc_path: 大纲文件路径
    :return: 去重后的全量文本字符串
    """
    if not os.path.exists(doc_path):
        raise FileNotFoundError(f"源文档不存在，请检查路径：{doc_path}")
    
    doc = Document(doc_path)
    full_text = []
    
    # 提取段落文本（跳过空段落）
    for paragraph in doc.paragraphs:
        para_text = paragraph.text.strip()
        if para_text:
            full_text.append(para_text)
    
    # 提取表格文本（跳过空单元格）
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    full_text.append(cell_text)
    
    # 段落级去重（解决大纲中课程简介等重复问题）
    unique_text = list(dict.fromkeys(full_text))
    return '\n'.join(unique_text)

# -------------------------- 2. 文本预处理：清洗冗余信息 --------------------------
def clean_syllabus_text(raw_text):
    """
    清洗文本：去除HTML标签、特殊符号、多余空格，保留核心内容
    :param raw_text: 从文档提取的原始文本
    :return: 清洗后的纯净文本
    """
    # 去除残留HTML标签
    text_no_html = re.sub(r'<[^>]+>', '', raw_text)
    # 保留中文、英文、数字、中文标点，删除其他符号
    text_clean = re.sub(r'[^\u4e00-\u9fa5a-zA-Z0-9，。；：！？、（）【】《》\s]', '', text_no_html)
    # 多余空格/换行转为单个空格
    text_clean = re.sub(r'[\n\r\s]+', ' ', text_clean)
    # 删除长度<2的无意义文本（如单个数字、字母）
    text_clean = ' '.join([seg for seg in text_clean.split() if len(seg) >= 2])
    return text_clean

# -------------------------- 3. 中文分词与停用词过滤 --------------------------
def segment_and_filter(clean_text):
    """
    基于jieba精确分词，过滤停用词，获取有效关键词及词频
    :param clean_text: 清洗后的文本
    :return: 过滤后关键词列表、词频统计Counter对象
    """
    # 停用词表（针对大纲场景，过滤无意义高频词）
    # 核心修改：删除停用词表中的"应用"，避免过滤"实际应用"
    stop_words = {
        '的','是','和','与','及','等','在','为','了','有','能','可','让','使','包括','包含',
        '课程','专业','学生','实验','教学','要求','掌握','熟悉','了解','熟练','培养','提高','构建',
        '开课','部门','学时','学分','类别','性质','适用','学期','代码','名称','英文','简介','说明',
        '日期','编写','审阅','教研室','学院','出版社','作者','版本','年','月','日','总','计','小计',
        '软件','用法','操作','步骤','必要','用于','可以','等等','实例','案例','分析','演练','习题','作业','以及','介绍','不仅',
        '问题'  # 去掉了"应用"
    }
    
    # 补充专业词汇，强化"实际应用"的分词优先级（保留并新增相关搭配）
    jieba.add_word('实际应用', freq=1000)  # 增加词频权重，确保优先分词
    jieba.add_word('单纯形法')
    jieba.add_word('蚁群优化算法')
    jieba.add_word('粒子群优化算法')
    jieba.add_word('模拟退火算法')
    jieba.add_word('线性规划')
    jieba.add_word('二次规划')
    jieba.add_word('最优化算法')
    jieba.add_word('工程应用')  # 新增相关应用场景词汇
    jieba.add_word('实际问题')
    
    # 精确分词
    words = jieba.lcut(clean_text, cut_all=False)
    
    # 过滤条件：
    # 1. 长度>1 2. 非停用词 3. 非纯数字 4. 过滤单独的"应用"（只保留组合词中的"应用"）
    filtered_words = []
    for word in words:
        if len(word) > 1 and word not in stop_words and not word.isdigit():
            if word != '应用':
                filtered_words.append(word)
    
    # 统计词频
    word_freq = Counter(filtered_words)
    
    return filtered_words, word_freq

# -------------------------- 4. 生成词云：基于关键词保存词云图片 --------------------------
def generate_syllabus_wordcloud(filtered_words, save_path="大纲词云.png"):
    # 拼接关键词为词云所需字符串
    word_str = ' '.join(filtered_words)
    
    # 词云配置（指定中文字体路径，避免乱码）
    wc = wordcloud.WordCloud(
        width=1000,
        height=700,
        background_color='white',
        font_path='C:/Windows/Fonts/SimHei.ttf',  # Windows系统黑体字体路径
        max_words=300,
        font_step=8,
        random_state=42,
        prefer_horizontal=0.8,
        relative_scaling=0.6,
        collocations=False  # 避免重复显示词组
    )
    
    # 生成并保存词云
    wc.generate(word_str)
    wc.to_file(save_path)
    print(f"词云已保存至：{os.path.abspath(save_path)}")

# -------------------------- 5. 保存实验结果：分词列表与词频统计 --------------------------
def save_experiment_result(filtered_words, word_freq, save_dir="实验结果"):

    # 新建文件夹（不存在则创建）
    if not os.path.exists(save_dir):
        os.makedirs(save_dir)
    
    # 保存分词结果
    with open(os.path.join(save_dir, "分词结果.txt"), 'w', encoding='utf-8') as f:
        f.write("《最优化算法》教学大纲中文分词结果（去停用词后）\n")
        f.write("="*50 + "\n")
        f.write(f"总关键词数：{len(filtered_words)}\n")
        f.write(f"不重复关键词数：{len(set(filtered_words))}\n")
        f.write("="*50 + "\n")
        f.write("分词列表（前100个）：\n")
        f.write(' '.join(filtered_words[:100]) + "\n")
    
    # 保存词频统计（Top50高频词）
    total_freq = sum(word_freq.values())
    with open(os.path.join(save_dir, "词频统计.txt"), 'w', encoding='utf-8') as f:
        f.write("《最优化算法》教学大纲关键词词频统计\n")
        f.write("="*50 + "\n")
        f.write(f"{'关键词':<15}{'词频':<8}{'占比（%）':<10}\n")
        f.write("="*50 + "\n")
        for word, freq in word_freq.most_common(50):
            ratio = (freq / total_freq) * 100
            f.write(f"{word:<15}{freq:<8}{ratio:.2f}\n")
    
    print(f"实验结果已保存至：{os.path.abspath(save_dir)}")

# -------------------------- 主函数--------------------------
if __name__ == "__main__":
    SYLLABUS_PATH = r"C:\Users\ahdezneytiry-boop1\Desktop\自然语言处理3\徐州医科大学理论课程教学大纲 最优化算法.docx"
    
    try:
        print("1. 正在读取教学大纲源文档...")
        raw_syllabus = read_syllabus_docx(SYLLABUS_PATH)
        print(f"成功读取文档，提取文本长度：{len(raw_syllabus)}字符")
        
        print("2. 正在清洗文本...")
        clean_text = clean_syllabus_text(raw_syllabus)
        print(f"文本清洗完成，纯净文本长度：{len(clean_text)}字符")
        
        print("3. 正在进行中文分词与停用词过滤...")
        filtered_words, word_freq = segment_and_filter(clean_text)
        print(f"分词完成，获得有效关键词：{len(filtered_words)}个")
        print(f"高频关键词Top10：{dict(word_freq.most_common(10))}")
        
        print("4. 正在生成词云...")
        generate_syllabus_wordcloud(filtered_words)
        
        print("5. 正在保存实验结果...")
        save_experiment_result(filtered_words, word_freq)
        
        print("\n全流程执行完成！")
    
    except Exception as e:
        print(f"执行出错：{str(e)}")